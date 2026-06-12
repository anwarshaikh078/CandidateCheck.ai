from io import BytesIO
from typing import List

import fitz
from docx import Document


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


class ResumeExtractionError(Exception):
    """Raised when resume text cannot be extracted."""


def extract_resume_text(uploaded_file) -> str:
    """Extract text from a Streamlit-uploaded PDF or DOCX file."""
    if uploaded_file is None:
        raise ResumeExtractionError("Please upload a PDF or DOCX resume.")

    filename = getattr(uploaded_file, "name", "") or ""
    extension = _extension(filename)
    if extension not in SUPPORTED_EXTENSIONS:
        raise ResumeExtractionError("Unsupported file type. Please upload a PDF or DOCX resume.")

    try:
        data = uploaded_file.getvalue()
        if extension == ".pdf":
            return _extract_pdf_text(data).strip()
        if extension == ".docx":
            return _extract_docx_text(data).strip()
    except ResumeExtractionError:
        raise
    except Exception as exc:
        raise ResumeExtractionError(
            "The resume could not be read. Please check that the file is not corrupted or password protected."
        ) from exc

    raise ResumeExtractionError("Unsupported file type. Please upload a PDF or DOCX resume.")


def extract_resume_links(uploaded_file) -> List[str]:
    """Extract embedded hyperlinks from a Streamlit-uploaded PDF or DOCX file."""
    if uploaded_file is None:
        return []

    filename = getattr(uploaded_file, "name", "") or ""
    extension = _extension(filename)
    if extension not in SUPPORTED_EXTENSIONS:
        return []

    try:
        data = uploaded_file.getvalue()
        if extension == ".pdf":
            return _extract_pdf_links(data)
        if extension == ".docx":
            return _extract_docx_links(data)
    except Exception:
        return []

    return []


def _extension(filename: str) -> str:
    dot_index = filename.rfind(".")
    return filename[dot_index:].lower() if dot_index != -1 else ""


def _extract_pdf_text(data: bytes) -> str:
    try:
        with fitz.open(stream=data, filetype="pdf") as document:
            text = "\n".join(page.get_text("text") for page in document)
    except Exception as exc:
        raise ResumeExtractionError(
            "The PDF could not be read. Please upload a text-based, unlocked PDF."
        ) from exc

    if not text.strip():
        raise ResumeExtractionError(
            "No readable text was found in the PDF. Scanned resumes may need OCR before analysis."
        )
    return text


def _extract_pdf_links(data: bytes) -> List[str]:
    links = []
    with fitz.open(stream=data, filetype="pdf") as document:
        for page in document:
            for link in page.get_links():
                uri = (link.get("uri") or "").strip()
                if uri:
                    links.append(uri)
    return _dedupe_links(links)


def _extract_docx_text(data: bytes) -> str:
    try:
        document = Document(BytesIO(data))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        table_cells = []
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        table_cells.append(cell.text)
        text = "\n".join(paragraphs + table_cells)
    except Exception as exc:
        raise ResumeExtractionError(
            "The DOCX could not be read. Please upload a valid Word document."
        ) from exc

    if not text.strip():
        raise ResumeExtractionError("No readable text was found in the DOCX resume.")
    return text


def _extract_docx_links(data: bytes) -> List[str]:
    document = Document(BytesIO(data))
    links = []
    for relationship in document.part.rels.values():
        if relationship.reltype.endswith("/hyperlink"):
            target = (relationship.target_ref or "").strip()
            if target:
                links.append(target)
    return _dedupe_links(links)


def _dedupe_links(links: List[str]) -> List[str]:
    seen = set()
    deduped = []
    for link in links:
        normalized = link.strip()
        key = normalized.lower()
        if normalized and key not in seen:
            seen.add(key)
            deduped.append(normalized)
    return deduped
